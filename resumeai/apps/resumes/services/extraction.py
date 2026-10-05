class ExtractionError(Exception):
    pass


def extract_pdf(path) -> tuple[str, int]:
    import pymupdf

    try:
        with pymupdf.open(path) as doc:
            pages = [page.get_text() for page in doc]
            return "\n".join(pages), len(pages)
    except Exception as exc:
        raise ExtractionError("Could not read the PDF (file may be corrupted or encrypted).") from exc


def extract_docx(path) -> tuple[str, int]:
    from docx import Document

    try:
        doc = Document(path)
    except Exception as exc:
        raise ExtractionError("Could not read the DOCX (file may be corrupted).") from exc
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells if c.text.strip()))
    return "\n".join(parts), 0


def extract_text(path, file_type: str) -> tuple[str, int]:
    text, pages = extract_pdf(path) if file_type == "pdf" else extract_docx(path)
    if not text.strip():
        raise ExtractionError("No text found. Scanned/image-only resumes are not supported yet.")
    return text, pages
